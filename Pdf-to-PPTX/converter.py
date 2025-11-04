from pdf2image import convert_from_path
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
import os

def convert_pdf_to_ppt(pdf_path, output_pptx):
    temp_folder = "temp_images"
    os.makedirs(temp_folder, exist_ok=True)

    pages = convert_from_path(pdf_path, dpi=300)
    first_page = pages[0]
    first_page_path = os.path.join(temp_folder, "page_1.png")
    first_page.save(first_page_path, "PNG")

    with Image.open(first_page_path) as im:
        img_width, img_height = im.size
        pdf_ratio = img_width / img_height

    prs = Presentation()
    base_width = Inches(13.33)
    prs.slide_width = base_width
    prs.slide_height = Inches(13.33 / pdf_ratio)

    for i, page in enumerate(pages):
        img_path = os.path.join(temp_folder, f"slide_{i+1}.png")
        page.save(img_path, "PNG")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(img_path, 0, 0, width=prs.slide_width, height=prs.slide_height)

    prs.save(output_pptx)
    for f in os.listdir(temp_folder):
        os.remove(os.path.join(temp_folder, f))
    os.rmdir(temp_folder)
